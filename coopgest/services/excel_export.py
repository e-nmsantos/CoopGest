"""
Excel report generation using openpyxl.
Produces a multi-sheet workbook for a cooperation project.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter

from coopgest.db import row_to_dict

# ── Colour constants ────────────────────────────────────────────────────────────
HEADER_BG = "1E40AF"       # blue
HEADER_FG = "FFFFFF"       # white
ROW_ALT_BG = "F3F4F6"     # light gray
ROW_MAIN_BG = "FFFFFF"    # white

# Tab colours for each sheet
SHEET_COLORS = {
    "Projeto": "1E40AF",
    "Quadro Lógico": "059669",
    "Beneficiários": "7C3AED",
    "Orçamento": "D97706",
    "Movimentos Financeiros": "DC2626",
    "Riscos": "B45309",
    "Tarefas": "0891B2",
    "Lições Aprendidas": "BE185D",
}


def _header_fill() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=HEADER_BG)


def _alt_fill() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=ROW_ALT_BG)


def _main_fill() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=ROW_MAIN_BG)


def _thin_border() -> Border:
    side = Side(style="thin", color="E5E7EB")
    return Border(left=side, right=side, top=side, bottom=side)


def _apply_header_row(ws, headers: list[str]) -> None:
    """Write bold blue header row with white text to row 1."""
    header_font = Font(bold=True, color=HEADER_FG, name="Calibri", size=10)
    header_fill = _header_fill()
    border = _thin_border()
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"


def _apply_data_rows(ws, rows: list[list[Any]], start_row: int = 2) -> None:
    """Write data rows with alternating background colours."""
    normal_font = Font(name="Calibri", size=9)
    border = _thin_border()
    for row_idx, row_data in enumerate(rows, start_row):
        fill = _alt_fill() if (row_idx % 2 == 0) else _main_fill()
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = normal_font
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _auto_column_widths(ws, headers: list[str]) -> None:
    """Set approximate column widths based on header length and data."""
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 40))
        ws.column_dimensions[col_letter].width = max(10, max_len + 4)


def _set_tab_color(ws, sheet_name: str) -> None:
    color = SHEET_COLORS.get(sheet_name, "1E40AF")
    ws.sheet_properties.tabColor = color


# ── Sheet builders ──────────────────────────────────────────────────────────────

def _build_projeto_sheet(wb: Workbook, projeto: dict) -> None:
    ws = wb.create_sheet("Projeto")
    _set_tab_color(ws, "Projeto")
    label_font = Font(bold=True, name="Calibri", size=10)
    value_font = Font(name="Calibri", size=10)
    border = _thin_border()

    fields = [
        ("Nome", projeto.get("nome", "")),
        ("Estado", projeto.get("estado", "")),
        ("Data de Início", projeto.get("data_inicio", "")),
        ("Data de Fim", projeto.get("data_fim", "")),
        ("Descrição", projeto.get("descricao", "")),
        ("Orçamento Total", projeto.get("orcamento_total", "")),
        ("País", projeto.get("pais", "")),
        ("Área Temática", projeto.get("area_tematica", "")),
    ]

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 50

    for row_idx, (label, value) in enumerate(fields, 1):
        lc = ws.cell(row=row_idx, column=1, value=label)
        lc.font = label_font
        lc.border = border
        lc.fill = _alt_fill() if row_idx % 2 == 0 else _main_fill()
        lc.alignment = Alignment(vertical="center")

        vc = ws.cell(row=row_idx, column=2, value=value)
        vc.font = value_font
        vc.border = border
        vc.fill = _alt_fill() if row_idx % 2 == 0 else _main_fill()
        vc.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row_idx].height = 18


def _build_quadro_logico_sheet(wb: Workbook, logframe: list[dict]) -> None:
    ws = wb.create_sheet("Quadro Lógico")
    _set_tab_color(ws, "Quadro Lógico")
    headers = [
        "Nível", "Resultado/Objetivo", "Indicador", "Unidade",
        "Baseline", "Meta", "Valor Atual", "Progresso (%)",
        "Estado", "Fonte de Verificação", "Responsável", "Frequência", "Pressupostos",
    ]
    _apply_header_row(ws, headers)
    rows = []
    for item in logframe:
        meta = float(item.get("meta") or 0)
        valor_atual = float(item.get("valor_atual") or 0)
        baseline = float(item.get("baseline") or 0)
        if meta > baseline:
            progresso = round(((valor_atual - baseline) / (meta - baseline)) * 100, 1)
        elif meta > 0:
            progresso = round((valor_atual / meta) * 100, 1)
        else:
            progresso = 0
        rows.append([
            item.get("nivel", ""),
            item.get("resultado", ""),
            item.get("indicador", ""),
            item.get("unidade", ""),
            item.get("baseline", ""),
            item.get("meta", ""),
            item.get("valor_atual", ""),
            progresso,
            item.get("estado", ""),
            item.get("fonte_verificacao", ""),
            item.get("responsavel_medicao", ""),
            item.get("frequencia_medicao", ""),
            item.get("pressupostos", ""),
        ])
    _apply_data_rows(ws, rows)
    _auto_column_widths(ws, headers)


def _build_beneficiarios_sheet(wb: Workbook, beneficiarios: list[dict]) -> None:
    ws = wb.create_sheet("Beneficiários")
    _set_tab_color(ws, "Beneficiários")
    headers = ["Nome", "Tipo", "Número", "Género", "Localização", "Data de Registo", "Notas"]
    _apply_header_row(ws, headers)
    rows = []
    for b in beneficiarios:
        rows.append([
            b.get("nome", ""),
            b.get("tipo", ""),
            b.get("numero", ""),
            b.get("genero", ""),
            b.get("localizacao", ""),
            b.get("data_registo", ""),
            b.get("descricao", ""),
        ])
    _apply_data_rows(ws, rows)
    _auto_column_widths(ws, headers)


def _build_orcamento_sheet(wb: Workbook, orcamento: list[dict]) -> None:
    ws = wb.create_sheet("Orçamento")
    _set_tab_color(ws, "Orçamento")
    headers = ["Tipo", "Categoria", "Descrição", "Valor Previsto", "Valor Real", "Execução (%)"]
    _apply_header_row(ws, headers)
    rows = []
    for o in orcamento:
        previsto = float(o.get("valor_previsto") or 0)
        real = float(o.get("valor_real") or 0)
        exec_pct = round((real / previsto) * 100, 1) if previsto else 0
        rows.append([
            o.get("tipo", ""),
            o.get("categoria", ""),
            o.get("descricao", ""),
            previsto,
            real,
            exec_pct,
        ])
    _apply_data_rows(ws, rows)
    _auto_column_widths(ws, headers)


def _build_movimentos_sheet(wb: Workbook, movimentos: list[dict]) -> None:
    ws = wb.create_sheet("Movimentos Financeiros")
    _set_tab_color(ws, "Movimentos Financeiros")
    headers = [
        "Data", "Tipo", "Categoria", "Descrição",
        "Valor", "Moeda", "Taxa Câmbio", "Entidade", "Estado",
    ]
    _apply_header_row(ws, headers)
    rows = []
    for m in movimentos:
        rows.append([
            m.get("data_movimento", ""),
            m.get("tipo", ""),
            m.get("categoria", ""),
            m.get("descricao", ""),
            m.get("valor", ""),
            m.get("moeda", ""),
            m.get("taxa_cambio", ""),
            m.get("entidade", ""),
            m.get("estado", ""),
        ])
    _apply_data_rows(ws, rows)
    _auto_column_widths(ws, headers)


def _build_riscos_sheet(wb: Workbook, riscos: list[dict]) -> None:
    ws = wb.create_sheet("Riscos")
    _set_tab_color(ws, "Riscos")
    headers = [
        "Descrição", "Probabilidade", "Impacto", "Nível de Risco",
        "Estado", "Mitigação", "Dono", "Próxima Revisão", "Plano de Contingência",
    ]
    _apply_header_row(ws, headers)
    rows = []
    prob_map = {"Baixo": 1, "Médio": 2, "Alto": 3}
    imp_map = {"Baixo": 1, "Médio": 2, "Alto": 3}
    nivel_map = {1: "Baixo", 2: "Médio", 3: "Alto", 4: "Muito Alto", 6: "Crítico", 9: "Crítico"}
    for r in riscos:
        prob = r.get("probabilidade", "Médio")
        imp = r.get("impacto", "Médio")
        nivel_score = prob_map.get(prob, 2) * imp_map.get(imp, 2)
        nivel = nivel_map.get(nivel_score, "Médio")
        rows.append([
            r.get("descricao", ""),
            prob,
            imp,
            nivel,
            r.get("estado", ""),
            r.get("mitigacao", ""),
            r.get("dono", ""),
            r.get("proxima_revisao", ""),
            r.get("plano_contingencia", ""),
        ])
    _apply_data_rows(ws, rows)
    _auto_column_widths(ws, headers)


def _build_tarefas_sheet(wb: Workbook, tarefas: list[dict]) -> None:
    ws = wb.create_sheet("Tarefas")
    _set_tab_color(ws, "Tarefas")
    headers = ["Título", "Estado", "Prioridade", "Responsável", "Data Início", "Data Fim"]
    _apply_header_row(ws, headers)
    rows = []
    for t in tarefas:
        rows.append([
            t.get("nome", ""),
            t.get("estado", ""),
            t.get("prioridade", ""),
            t.get("responsavel", ""),
            t.get("data_inicio", ""),
            t.get("data_fim", ""),
        ])
    _apply_data_rows(ws, rows)
    _auto_column_widths(ws, headers)


def _build_licoes_sheet(wb: Workbook, licoes: list[dict]) -> None:
    ws = wb.create_sheet("Lições Aprendidas")
    _set_tab_color(ws, "Lições Aprendidas")
    headers = [
        "Título", "Tipo", "Área", "Fase", "Impacto",
        "Descrição", "Recomendação", "Criado Por",
    ]
    _apply_header_row(ws, headers)
    rows = []
    for l in licoes:
        rows.append([
            l.get("titulo", ""),
            l.get("tipo", ""),
            l.get("area", ""),
            l.get("fase_projeto", ""),
            l.get("impacto", ""),
            l.get("descricao", ""),
            l.get("recomendacao", ""),
            l.get("criado_por", ""),
        ])
    _apply_data_rows(ws, rows)
    _auto_column_widths(ws, headers)


# ── Public API ──────────────────────────────────────────────────────────────────

def generate_project_excel(conn, projeto_id: int) -> bytes:
    """
    Generate a multi-sheet Excel workbook for a project and return the bytes.
    """
    projeto_row = conn.execute("SELECT * FROM projetos WHERE id=?", (projeto_id,)).fetchone()
    projeto = row_to_dict(projeto_row) or {}

    logframe = [
        row_to_dict(r)
        for r in conn.execute(
            "SELECT * FROM impacto_quadro_logico WHERE projeto_id=? ORDER BY nivel, id",
            (projeto_id,),
        ).fetchall()
    ]
    beneficiarios = [
        row_to_dict(r)
        for r in conn.execute(
            "SELECT * FROM beneficiarios WHERE projeto_id=? AND arquivado=0 ORDER BY nome",
            (projeto_id,),
        ).fetchall()
    ]
    orcamento = [
        row_to_dict(r)
        for r in conn.execute(
            "SELECT * FROM orcamento WHERE projeto_id=? ORDER BY tipo, categoria",
            (projeto_id,),
        ).fetchall()
    ]
    movimentos = [
        row_to_dict(r)
        for r in conn.execute(
            "SELECT * FROM movimentos_financeiros WHERE projeto_id=? ORDER BY data_movimento DESC",
            (projeto_id,),
        ).fetchall()
    ]
    riscos = [
        row_to_dict(r)
        for r in conn.execute(
            "SELECT * FROM riscos WHERE projeto_id=? ORDER BY criado_em",
            (projeto_id,),
        ).fetchall()
    ]
    tarefas = [
        row_to_dict(r)
        for r in conn.execute(
            "SELECT * FROM tarefas WHERE projeto_id=? ORDER BY estado, data_fim",
            (projeto_id,),
        ).fetchall()
    ]
    licoes = [
        row_to_dict(r)
        for r in conn.execute(
            "SELECT * FROM licoes_aprendidas WHERE projeto_id=? ORDER BY criado_em DESC",
            (projeto_id,),
        ).fetchall()
    ]

    wb = Workbook()
    # Remove the default sheet created by openpyxl
    default_sheet = wb.active
    wb.remove(default_sheet)

    _build_projeto_sheet(wb, projeto)
    _build_quadro_logico_sheet(wb, logframe)
    _build_beneficiarios_sheet(wb, beneficiarios)
    _build_orcamento_sheet(wb, orcamento)
    _build_movimentos_sheet(wb, movimentos)
    _build_riscos_sheet(wb, riscos)
    _build_tarefas_sheet(wb, tarefas)
    _build_licoes_sheet(wb, licoes)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
